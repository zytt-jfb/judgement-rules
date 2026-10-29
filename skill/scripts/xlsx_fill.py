#!/usr/bin/env python
"""按列定义把 CSV 写成带样式的 Excel。

适合"客户给了个模板，要把结果填进去"的场景。列定义用 JSON 描述，能覆盖：
列顺序与宽度、取值映射、分级配色、冻结窗格、自动筛选、按内容估算行高。

用法：
    python xlsx_fill.py --csv rows.csv --spec columns.json --out 结果.xlsx

columns.json 结构：
{
  "sheet": "Sheet1",
  "columns": [
    {"header": "序号", "source": "__index__", "width": 6, "align": "center"},
    {"header": "姓名", "source": "姓名", "width": 13, "align": "center"},
    {"header": "问题分类", "source": "问题分类", "width": 22},
    {"header": "检查人员", "source": null, "width": 10}
  ],
  "color_by": {
    "column": "问题分类",
    "rules": [
      {"contains": ["非本人", "过期", "状态异常"], "fill": "FFC7CE"},
      {"equals": "无问题", "fill": "E2EFDA"},
      {"startswith": "未核验", "fill": "EDEDED"},
      {"default": "FFE0B2"}
    ]
  },
  "row_height": {"by": "问题备注说明", "chars_per_line": 38, "line_height": 14, "pad": 6},
  "freeze": "A2",
  "autofilter": true,
  "font": {"name": "微软雅黑", "size": 10}
}

source 为 null 表示留空列（留给人工填写）。
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def match_rule(value: str, rules: list[dict]) -> str | None:
    for r in rules:
        if "equals" in r and value == r["equals"]:
            return r.get("fill")
        if "startswith" in r and value.startswith(r["startswith"]):
            return r.get("fill")
        if "contains" in r and any(k in value for k in r["contains"]):
            return r.get("fill")
    for r in rules:
        if "default" in r:
            return r["default"]
    return None


def main() -> int:
    ap = argparse.ArgumentParser(description="按列定义把 CSV 写成带样式的 Excel")
    ap.add_argument("--csv", required=True)
    ap.add_argument("--spec", required=True, help="列定义 JSON")
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    rows = list(csv.DictReader(pathlib.Path(args.csv).open(encoding="utf-8-sig", newline="")))
    spec = json.loads(pathlib.Path(args.spec).read_text(encoding="utf-8"))
    cols = spec["columns"]

    font_cfg = spec.get("font", {})
    fname = font_cfg.get("name", "微软雅黑")
    fsize = font_cfg.get("size", 10)

    wb = Workbook()
    ws = wb.active
    ws.title = spec.get("sheet", "Sheet1")
    ws.append([c["header"] for c in cols])

    for i, row in enumerate(rows, 1):
        line = []
        for c in cols:
            src = c.get("source")
            if src == "__index__":
                line.append(i)
            elif src is None:
                line.append("")
            else:
                line.append(row.get(src, ""))
        ws.append(line)

    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = c.get("width", 14)

    for cell in ws[1]:
        cell.font = Font(name=fname, size=fsize, bold=True)
        cell.fill = PatternFill("solid", fgColor=spec.get("header_fill", "D9E1F2"))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30

    color_by = spec.get("color_by")
    color_idx = None
    if color_by:
        for i, c in enumerate(cols, 1):
            if c["header"] == color_by["column"]:
                color_idx = i
                break

    rh = spec.get("row_height")
    rh_idx = None
    if rh:
        for i, c in enumerate(cols, 1):
            if c["header"] == rh["by"]:
                rh_idx = i
                break

    for r in range(2, ws.max_row + 1):
        for ci, c in enumerate(cols, 1):
            cell = ws.cell(row=r, column=ci)
            cell.font = Font(name=fname, size=fsize)
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=c.get("wrap", c.get("width", 14) < 60),
                horizontal="center" if c.get("align", "left") == "center" else "left",
            )
        if color_idx:
            v = str(ws.cell(row=r, column=color_idx).value or "")
            fill = match_rule(v, color_by.get("rules", []))
            if fill:
                ws.cell(row=r, column=color_idx).fill = PatternFill("solid", fgColor=fill)
        if rh_idx:
            text = str(ws.cell(row=r, column=rh_idx).value or "")
            n = max(math.ceil(len(text) / rh.get("chars_per_line", 38)), 1)
            ws.row_dimensions[r].height = max(20, rh.get("line_height", 14) * n + rh.get("pad", 6))

    if spec.get("freeze"):
        ws.freeze_panes = spec["freeze"]
    if spec.get("autofilter", True):
        ws.auto_filter.ref = ws.dimensions

    out = pathlib.Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"已生成 {out}（{len(rows)} 行 × {len(cols)} 列）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
